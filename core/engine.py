"""Blockly 워크스페이스(JSON)를 해석해서 실제 키보드/마우스 동작으로 실행하는 엔진."""
import datetime
import os
import time
import webbrowser
from urllib.parse import quote, urlsplit

import pyperclip
from pynput.keyboard import Controller as KbController, Key

from . import datetime_blocks as dtb
from .recorder import replay

kb = KbController()

SPECIAL_KEYS = {
    "enter": Key.enter,
    "tab": Key.tab,
    "delete": Key.delete,
    "backspace": Key.backspace,
    "down": Key.down,
    "up": Key.up,
    "left": Key.left,
    "right": Key.right,
}


def read_rows(path, sheet, start):
    """엑셀 파일(.xlsx/.xlsm/.xls)에서 start 행부터 [(행번호, [셀값...]), ...] 를 읽는다. 완전히 빈 줄은 제외."""
    if path.lower().endswith(".xls"):
        import xlrd
        try:
            wb = xlrd.open_workbook(path)
        except Exception as e:
            raise EngineError(f"엑셀 파일을 열 수 없습니다: {e}")
        # 시트 이름이 없거나 맞지 않으면 첫 번째 시트를 쓴다
        ws = wb.sheet_by_name(sheet) if sheet in wb.sheet_names() else wb.sheet_by_index(0)
        rows = []
        for i in range(start - 1, ws.nrows):
            row = []
            for c in ws.row(i):
                if c.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    row.append(None)
                elif c.ctype == xlrd.XL_CELL_DATE:
                    row.append(xlrd.xldate.xldate_as_datetime(c.value, wb.datemode))
                elif c.ctype == xlrd.XL_CELL_BOOLEAN:
                    row.append(bool(c.value))
                else:
                    row.append(c.value)
            rows.append((i + 1, row))
    else:
        from openpyxl import load_workbook
        try:
            wb = load_workbook(path, read_only=True, data_only=True)
        except Exception as e:
            raise EngineError(f"엑셀 파일을 열 수 없습니다: {e}")
        try:
            # 시트 이름이 없거나 맞지 않으면 첫 번째 시트를 쓴다
            ws = wb[sheet] if sheet in wb.sheetnames else wb.worksheets[0]
            ws.reset_dimensions()  # 일부 프로그램이 만든 파일은 범위가 A1 로만 적혀 있어 나머지 줄을 못 읽는다
            rows = [(n, list(r)) for n, r in enumerate(ws.iter_rows(min_row=start, values_only=True), start)]
        finally:
            wb.close()
    return [(n, r) for n, r in rows if any(c not in (None, "") for c in r)]


class Stopped(Exception):
    pass


class EngineError(Exception):
    pass


def fmt(v):
    """엑셀 셀 값을 입력용 문자열로 변환."""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, datetime.datetime):
        return v.strftime("%Y-%m-%d")
    return str(v)


def paste_text(text):
    """한글도 깨지지 않도록 클립보드 + Ctrl+V 로 입력."""
    try:
        old = pyperclip.paste()
    except Exception:
        old = None
    pyperclip.copy(text)
    time.sleep(0.05)
    with kb.pressed(Key.ctrl):
        kb.press("v")
        kb.release("v")
    time.sleep(0.12)
    if old is not None:
        pyperclip.copy(old)


class Engine:
    STATEMENTS = {"if_else", "open_website", "open_program", "excel_each", "repeat_times", "repeat_forever", "type_text", "key_press", "wait", "macro_play"}

    def __init__(self, stop_event, on_step=None):
        self.stop = stop_event
        self.on_step = on_step or (lambda _id: None)
        self.row = None      # 엑셀 블록 안에서 현재 줄의 값들
        self.hol = dtb.Holidays()  # 휴일 설정 블록 내용 (없으면 토/일만 휴무)
        self.col = 0         # 엑셀 블록에서 정한 기본(가져올) 열 번호(0부터)
        self.row_no = 0

    # ---- 실행 ----
    def run(self, state):
        tops = state.get("blocks", {}).get("blocks", [])
        tops = [b for b in tops if b["type"] in self.STATEMENTS]
        if not tops:
            raise EngineError("실행할 블록이 없습니다.")
        self.load_holidays(state)
        for b in sorted(tops, key=lambda b: b.get("y", 0)):
            self.exec_chain(b)

    def load_holidays(self, state):
        """'휴일 설정' 블록(연결 없이 아무 곳에나 놓는 블록)을 모두 합쳐서 읽는다."""
        blocks = [b for b in state.get("blocks", {}).get("blocks", []) if b["type"] == "holiday_settings"]
        if not blocks:
            return
        sat = sun = True
        dates = set()
        for b in blocks:
            f = b.get("fields", {})
            sat, sun = bool(f.get("SAT", True)), bool(f.get("SUN", True))
            try:
                dates |= dtb.parse_holidays(f.get("HOLIDAYS", ""))  # 예전 저장 파일의 글자 입력칸
            except dtb.DateError as e:
                raise EngineError(str(e))
            dates |= self.holiday_dates(b)
        self.hol = dtb.Holidays(sat, sun, dates)

    @staticmethod
    def holiday_dates(settings):
        """'휴일 설정' 블록 안에 연결된 [휴일 날짜] / [휴일 기간] 블록들을 날짜 집합으로 변환."""
        from datetime import date, timedelta

        def ymd(f, prefix=""):
            y, m, d = (int(float(f.get(prefix + k, 0))) for k in ("YEAR", "MONTH", "DAY"))
            try:
                return date(y, m, d)
            except ValueError:
                raise EngineError(f"휴일 설정의 날짜가 올바르지 않습니다: {y}년 {m}월 {d}일")

        result = set()
        blk = settings.get("inputs", {}).get("DATES", {}).get("block")
        while blk:
            f = blk.get("fields", {})
            if blk["type"] == "holiday_date":
                result.add(ymd(f))
            elif blk["type"] == "holiday_range":
                a, b = ymd(f, "S_"), ymd(f, "E_")
                if b < a:
                    raise EngineError(f"휴일 기간의 끝 날짜가 시작 날짜보다 빠릅니다: {a} ~ {b}")
                if (b - a).days > 3660:
                    raise EngineError("휴일 기간이 너무 깁니다. (최대 10년)")
                result.update(a + timedelta(days=i) for i in range((b - a).days + 1))
            blk = blk.get("next", {}).get("block")
        return result

    @staticmethod
    def adjust_chain(block, name):
        """날짜/시간 블록 오른쪽에 가로로 이어 붙인 조정 블록들의 fields 를 왼쪽부터 순서대로 모은다."""
        chain = []
        nxt = block.get("inputs", {}).get(name, {}).get("block")
        while nxt:
            chain.append(nxt.get("fields", {}))
            nxt = nxt.get("inputs", {}).get("NEXT", {}).get("block")
        return chain

    def check(self):
        if self.stop.is_set():
            raise Stopped()

    def sleep(self, sec):
        end = time.time() + sec
        while time.time() < end:
            self.check()
            time.sleep(min(0.05, max(0, end - time.time())))

    def exec_chain(self, block):
        while block:
            self.check()
            self.on_step(block.get("id"))
            getattr(self, "do_" + block["type"])(block)
            block = block.get("next", {}).get("block")

    # ---- 값 블록 ----
    def eval(self, block, what="입력칸"):
        if not block:
            raise EngineError(f"{what}이 비어 있습니다. 값 블록(텍스트 등)을 끼워 주세요.")
        t = block["type"]
        f = block.get("fields", {})
        if t == "text":
            return f.get("TEXT", "")
        if t == "math_number":
            return fmt(f.get("NUM", 0))
        if t == "excel_value":
            if self.row is None:
                raise EngineError("'엑셀 값' 블록은 '엑셀의 각 줄마다 실행' 블록 안에서만 쓸 수 있습니다.")
            v = self.row[self.col] if self.col < len(self.row) else None
            return "" if v is None else fmt(v)
        if t == "date_value":
            return dtb.eval_date(f, self.hol, chain=self.adjust_chain(block, "ADJ"))
        if t == "time_value":
            return dtb.eval_time(f, chain=self.adjust_chain(block, "ADJ"))
        if t == "excel_row_no":
            if self.row is None:
                raise EngineError("'엑셀 줄 번호' 블록은 '엑셀의 각 줄마다 실행' 블록 안에서만 쓸 수 있습니다.")
            return str(self.row_no)
        if t == "text_join":
            n = block.get("extraState", {}).get("itemCount", 2)
            parts = []
            for i in range(n):
                inp = block.get("inputs", {}).get(f"ADD{i}", {}).get("block")
                if inp:
                    parts.append(self.eval(inp))
            return "".join(parts)
        raise EngineError(f"지원하지 않는 값 블록: {t}")

    @staticmethod
    def col_index(letters):
        from openpyxl.utils import column_index_from_string
        try:
            return column_index_from_string(str(letters).strip().upper()) - 1
        except ValueError:
            raise EngineError(f"열 이름이 올바르지 않습니다: {letters}")

    # ---- 문장 블록 ----
    def do_excel_each(self, b):
        f = b["fields"]
        path = f.get("FILE", "").strip()
        if not path:
            raise EngineError("엑셀 블록에 파일이 지정되지 않았습니다.")
        start = int(f.get("START", 1))
        col = self.col_index(f.get("COL", "A"))
        rows = read_rows(path, f.get("SHEET", "").strip(), start)
        body = b.get("inputs", {}).get("DO", {}).get("block")
        saved = (self.row, self.row_no, self.col)
        self.col = col
        try:
            for n, r in rows:
                self.check()
                self.row, self.row_no = r, n
                self.on_step(b.get("id"))
                self.exec_chain(body)
        finally:
            self.row, self.row_no, self.col = saved

    def do_repeat_times(self, b):
        body = b.get("inputs", {}).get("DO", {}).get("block")
        for _ in range(int(b["fields"].get("TIMES", 1))):
            self.check()
            self.on_step(b.get("id"))
            self.exec_chain(body)

    def do_repeat_forever(self, b):
        body = b.get("inputs", {}).get("DO", {}).get("block")
        if not body:
            raise EngineError("계속 반복 블록 안이 비어 있습니다.")
        while True:
            self.check()
            self.on_step(b.get("id"))
            self.exec_chain(body)

    @staticmethod
    def input_block(block, name):
        """입력칸에 끼운 블록. 사용자가 끼운 블록이 없으면 기본으로 들어 있는 흐린(shadow) 블록을 쓴다."""
        inp = block.get("inputs", {}).get(name, {})
        return inp.get("block") or inp.get("shadow")

    # ---- 조건(if) ----
    @staticmethod
    def _num(s):
        try:
            return float(str(s).replace(",", ""))
        except ValueError:
            return None

    def eval_bool(self, block, what="조건칸"):
        if not block:
            raise EngineError(f"{what}이 비어 있습니다. 조건 블록을 끼워 주세요.")
        t = block["type"]
        f = block.get("fields", {})
        if t == "cond_bday":
            return dtb.is_business_day(f, self.hol, chain=self.adjust_chain(block, "ADJ"))
        if t == "cond_not":
            return not self.eval_bool(self.input_block(block, "A"))
        if t == "cond_logic":
            a = self.eval_bool(self.input_block(block, "A"))
            if f.get("OP") == "OR":
                return a or self.eval_bool(self.input_block(block, "B"))
            return a and self.eval_bool(self.input_block(block, "B"))
        if t == "cond_empty":
            return self.eval(self.input_block(block, "A"), "값 입력칸").strip() == ""
        if t == "cond_compare":
            a = self.eval(self.input_block(block, "A"), "비교할 값").strip()
            b = self.eval(self.input_block(block, "B"), "비교할 값").strip()
            na, nb = self._num(a), self._num(b)
            if na is not None and nb is not None:  # 둘 다 숫자면 숫자로 비교 (12 와 12.0 은 같음)
                a, b = na, nb
            op = f.get("OP", "EQ")
            if op in ("EQ", "NEQ"):
                return (a == b) == (op == "EQ")
            return {"GT": a > b, "GTE": a >= b, "LT": a < b, "LTE": a <= b}[op]
        raise EngineError(f"지원하지 않는 조건 블록: {t}")

    def do_if_else(self, b):
        cond = self.eval_bool(self.input_block(b, "COND"))
        body = b.get("inputs", {}).get("DO" if cond else "ELSE", {}).get("block")
        self.exec_chain(body)

    def do_holiday_date(self, b):  # '휴일 설정' 안에서만 의미 있음 (설정은 시작할 때 읽음)
        pass

    do_holiday_range = do_holiday_date

    def do_open_website(self, b):
        url = self.eval(self.input_block(b, "URL"), "웹사이트 주소 입력칸").strip()
        if not url or url.lower() in ("http://", "https://"):
            raise EngineError("웹사이트 주소가 비어 있습니다.")
        if "://" not in url:
            url = "https://" + url  # 주소만 적어도 열리도록 https 를 붙임
        parts = urlsplit(url)
        if parts.scheme.lower() not in ("http", "https") or not parts.netloc:
            raise EngineError(f"http 또는 https 주소만 열 수 있습니다: {url}")
        url = quote(url, safe=":/?#[]@!$&'()*+,;=%~-._")  # 공백 등을 안전하게 변환
        if not webbrowser.open(url, new=2):  # 기본 브라우저, 새 탭
            raise EngineError("기본 브라우저를 열 수 없습니다.")
        self.sleep(float(b.get("fields", {}).get("WAIT", 3)))  # 페이지가 뜰 때까지 기다림

    def do_open_program(self, b):
        f = b["fields"]
        path = f.get("PATH", "").strip().strip('"')
        if not path:
            raise EngineError("프로그램 열기 블록에 실행 파일이 지정되지 않았습니다.")
        if not os.path.isfile(path):
            raise EngineError(f"파일을 찾을 수 없습니다: {path}")
        try:
            os.startfile(path, cwd=os.path.dirname(path))  # exe, bat, lnk 등을 더블클릭한 것처럼 실행
        except OSError as e:
            raise EngineError(f"프로그램을 열 수 없습니다: {e}")
        self.sleep(float(f.get("WAIT", 3)))  # 프로그램 창이 뜰 때까지 기다림

    def do_type_text(self, b):
        text = self.eval(b.get("inputs", {}).get("TEXT", {}).get("block"), "타이핑 입력칸")
        paste_text(text)

    def do_key_press(self, b):
        key = SPECIAL_KEYS[b["fields"].get("KEY", "enter")]
        kb.press(key)
        kb.release(key)
        time.sleep(0.05)

    def do_wait(self, b):
        self.sleep(float(b["fields"].get("SEC", 1)))

    def do_macro_play(self, b):
        events = b.get("extraState", {}).get("events", [])
        if not events:
            raise EngineError("녹화 블록에 녹화된 동작이 없습니다. 녹화 버튼을 눌러 주세요.")
        replay(events, self.stop, Stopped)
